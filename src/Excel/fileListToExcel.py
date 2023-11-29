import os
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import sys
from datetime import datetime

def add_folder_info(ws, folder_path, row_num):
    for root, dirs, files in os.walk(folder_path):
        if files:  # Only process directories that contain files
            folder_structure = os.path.relpath(root, folder_path)
            folder_structure_list = folder_structure.split(os.path.sep)

            for name in files:
                # If the file is directly under the input folder, do not include the folder name
                if folder_structure == '.':
                    row_data = [name]
                else:
                    row_data = folder_structure_list + [name]

                for i, folder_name in enumerate(row_data):
                    # 파일/폴더 이름에 하이퍼링크 추가
                    hyperlink = os.path.abspath(os.path.join(root, name))
                    ws.cell(row=row_num, column=i+1, value=folder_name).hyperlink = hyperlink

                row_num += 1

def create_excel_file(folder_path):
    # 부모 폴더 경로 추출
    parent_folder_path = os.path.dirname(folder_path)

    # 파일 이름에 타임스탬프 추가
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")

    # 새로운 Excel 워크북 생성
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "File list to Excel"

    # 폴더 구조를 순회하며 Excel 시트에 정보 추가
    row_num = 1
    add_folder_info(ws, folder_path, row_num)

    # 부모 폴더에 타임스탬프가 추가된 Excel 파일 저장
    excel_file_name = f"FileList_{timestamp}.xlsx"
    excel_file_path = os.path.join(parent_folder_path, excel_file_name)
    wb.save(excel_file_path)

    return excel_file_path

def open_file(file_path):
    system = sys.platform
    if system.startswith("win"):
        os.system(f'start "" "{file_path}"')
    elif system == "darwin":  # macOS
        os.system(f'open "{file_path}"')
    else:
        print("Windows와 macOS만 지원합니다.")

def main():
    print("특정 폴더 아래에 있는 파일과 폴더 리스트를 Excel로 출력해주는 프로그램입니다.")
    print("폴더의 절대 경로를 입력하세요: ")
    folder_path = input()

    if not os.path.exists(folder_path):
        print("지정한 폴더 경로가 존재하지 않습니다.")
        return

    excel_file_path = create_excel_file(folder_path)
    print(f"Excel 파일이 생성되었습니다: {excel_file_path}")

    open_file(excel_file_path)

if __name__ == "__main__":
    main()